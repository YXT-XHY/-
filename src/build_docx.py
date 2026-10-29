# -*- coding: utf-8 -*-
"""Build the Chinese Word report from report.md with python-docx."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "report.md"
OUT = ROOT / "deliverables" / "声纹帕金森早期筛查_项目调研与架构评估报告.docx"
OUT.parent.mkdir(exist_ok=True)

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
LATIN_FONT = "Times New Roman"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
INK = RGBColor(0x17, 0x20, 0x33)
MUTED = RGBColor(0x5B, 0x67, 0x7A)
ACCENT = RGBColor(0x17, 0x7E, 0x73)
ORANGE = RGBColor(0xC4, 0x5B, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

USABLE_CM = 15.9


# ----------------------------------------------------------------------------- basic helpers
def set_run_font(run, name=BODY_FONT, size=10.5, bold=None, color=None, italic=None):
    run.font.name = LATIN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_run(p, text, name=BODY_FONT, size=10.5, bold=None, color=None, italic=None):
    r = p.add_run(text)
    set_run_font(r, name, size, bold, color, italic)
    return r


def add_hyperlink(p, url, text, size=10.5):
    part = p.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), LATIN_FONT)
    rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    rFonts.set(qn("w:eastAsia"), BODY_FONT)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    p._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_inline_runs(p, text, size=10.5, color=INK, base_bold=None):
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            add_run(p, text[pos:m.start()], size=size, color=color, bold=base_bold)
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            add_run(p, token[2:-2], size=size, color=color, bold=True)
        elif token.startswith("[") and "](" in token:
            mm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if mm:
                add_hyperlink(p, mm.group(2), mm.group(1), size=size)
            else:
                add_run(p, token, size=size, color=color, bold=base_bold)
        elif token.startswith("*") and token.endswith("*"):
            add_run(p, token[1:-1], size=size, color=color, bold=base_bold, italic=True)
        else:
            add_run(p, token, size=size, color=color, bold=base_bold)
        pos = m.end()
    if pos < len(text):
        add_run(p, text[pos:], size=size, color=color, bold=base_bold)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=40, start=80, bottom=40, end=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def _replace_tbl_borders(tbl, borders):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    tblPr.insert_element_before(
        borders,
        "w:shd",
        "w:tblLayout",
        "w:tblCellMar",
        "w:tblLook",
        "w:tblCaption",
        "w:tblDescription",
    )


def _replace_tc_borders(cell, tcBorders):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    tcPr.insert_element_before(
        tcBorders,
        "w:shd",
        "w:noWrap",
        "w:tcMar",
        "w:textDirection",
        "w:tcFitText",
        "w:vAlign",
        "w:hideMark",
    )


def set_table_borders(tbl, color="BFC9D9", sz=4):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    _replace_tbl_borders(tbl, borders)


def set_table_no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)


def set_cell_no_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)


def set_table_top_border(tbl, color="CBD5E1", sz=4):
    borders = OxmlElement("w:tblBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(sz))
    top.set(qn("w:space"), "0")
    top.set(qn("w:color"), color)
    borders.append(top)
    for edge in ("left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    _replace_tbl_borders(tbl, borders)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    trPr.append(cantSplit)


def add_bottom_border(p, color="1F3A5F", sz=24):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_top_border(p, color="CBD5E1", sz=4):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(sz))
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), color)
    pBdr.append(top)
    pPr.append(pBdr)


def add_field(p, instr, placeholder="1", size=8.5, color=MUTED):
    run = p.add_run()
    set_run_font(run, size=size, color=color)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = placeholder
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(it)
    run._r.append(fld2)
    run._r.append(t)
    run._r.append(fld3)


# ----------------------------------------------------------------------------- document setup
def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.55)
    section.right_margin = Cm(2.55)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.1)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = LATIN_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.42
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.space_before = Pt(0)

    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        st = styles[f"Heading {lvl}"]
        st.font.name = HEAD_FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(14 if lvl == 1 else 8)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15

    # header (not first page)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(p, "声纹帕金森早期筛查 · 项目调研与架构评估", size=8, color=MUTED)
    add_bottom_border(p, color="CBD5E1", sz=4)

    # footer: one paragraph with a right tab stop (avoids table grid lines)
    footer = section.footer
    ft = footer.paragraphs[0]
    ft.clear()
    ft.alignment = WD_ALIGN_PARAGRAPH.LEFT
    ft.paragraph_format.tab_stops.clear_all()
    ft.paragraph_format.tab_stops.add_tab_stop(Cm(8.255), WD_TAB_ALIGNMENT.CLEAR)
    ft.paragraph_format.tab_stops.add_tab_stop(Cm(15.6), WD_TAB_ALIGNMENT.RIGHT)
    add_run(ft, "内部资料 · 请勿外传", size=8, color=MUTED)
    add_run(ft, "\t第 ", size=8, color=MUTED)
    add_field(ft, "PAGE", "1")
    add_run(ft, " 页 / 共 ", size=8, color=MUTED)
    add_field(ft, "NUMPAGES", "1")
    add_run(ft, " 页", size=8, color=MUTED)
    add_top_border(ft, color="CBD5E1", sz=4)

    # first page header/footer empty
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""

    props = doc.core_properties
    props.title = "声纹驱动的帕金森病早期筛查：项目调研、论文证据与架构评估"
    props.subject = "帕金森病语音早筛项目立项建议书"
    props.author = "项目筹备组"
    props.keywords = "帕金森病; 声纹; 语音; 早期筛查; 架构评估; 数字生物标志物"
    props.comments = "内部研究资料，不构成医疗、法律或投资建议。"
    return doc


# ----------------------------------------------------------------------------- cover & navigation
def cover(doc):
    # top band
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Cm(USABLE_CM)
    shade_cell(cell, "1F3A5F")
    set_table_no_borders(tbl)
    cell.paragraphs[0].paragraph_format.space_before = Pt(0)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    r = cell.paragraphs[0].add_run(" ")
    r.font.size = Pt(2)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    add_run(p, "项目立项建议书", size=11, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "声纹驱动的帕金森病", size=30, bold=True, color=NAVY, name=HEAD_FONT)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_after = Pt(0)
    add_run(p, "早期筛查", size=30, bold=True, color=NAVY, name=HEAD_FONT)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    add_run(p, "项目调研、论文证据与架构评估", size=16, color=INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    add_run(p, "v1.0 · 供立项评审", size=11.5, color=ACCENT)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    add_bottom_border(p, color="1F3A5F", sz=18)

    # metadata
    meta = [
        ("编制单位", "项目筹备组"),
        ("编制日期", "2026 年 10 月 27 日"),
        ("版本", "v1.0（供立项评审）"),
        ("密级", "内部资料"),
        ("检索截止", "2026 年 10 月 27 日"),
        ("报告性质", "快速证据扫描 + 架构评估（非系统综述）"),
    ]
    for k, v in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.2), WD_TAB_ALIGNMENT.LEFT)
        add_run(p, k, size=9, color=MUTED)
        add_run(p, "\t", size=9, color=INK)
        add_run(p, v, size=9, color=INK)

    # bottom note
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(58)
    add_top_border(p, color="CBD5E1", sz=4)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    add_run(p, "本报告用于内部立项讨论，不构成医疗、法律或投资建议。引用与数据来源截至 2026-10-27；性能数字均为原文报告，未做复现。", size=8, color=MUTED)

    doc.add_page_break()


def navigation(doc):
    h = doc.add_heading("阅读导航", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    add_run(p, "本报告共 8 章 + 附录。建议先读“摘要”和“8.3 需要决策的事项”，再按角色跳读。", size=10.5, color=MUTED)

    rows = [
        ("摘要", "一页看懂结论与决策请求", "核心结论、五个发现、请求决策"),
        ("1 项目定义与边界", "做什么、不做什么、成功标准", "预期用途、目标人群、Go/No-Go"),
        ("2 论文与证据扫描", "语音标志、关键研究、真实世界性能", "性能掉点、系统综述、证据缺口"),
        ("3 数据集与开源项目", "可用数据、可复用代码、许可风险", "数据集对比、GitHub 分层、许可矩阵"),
        ("4 参考架构", "九层架构、双通道、端云协同", "架构图、数据流、部署形态、NFR"),
        ("5 架构评估", "五方案评分、ADR、风险矩阵", "加权评分、关键决策、综合评级"),
        ("6 临床、监管与合规", "研究设计、PPV、NMPA、PIPL", "样本量、监管路径、伦理"),
        ("7 实施路线图", "阶段、团队、预算、Go/No-Go", "30 个月排期、预算量级、KPI"),
        ("8 结论与建议", "总体判断与决策事项", "三条建议、局限"),
        ("附录 / 参考文献", "公式、术语、出处", "36 条参考文献"),
    ]
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)
    widths = [Cm(4.2), Cm(6.2), Cm(5.5)]
    hdr = ["章节", "核心问题", "关键产出"]
    for i, text in enumerate(hdr):
        cell = tbl.rows[0].cells[i]
        cell.width = widths[i]
        shade_cell(cell, "1F3A5F")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, text, size=9, bold=True, color=WHITE)
    set_repeat_table_header(tbl.rows[0])
    for j, row in enumerate(rows):
        cells = tbl.add_row().cells
        prevent_row_split(tbl.rows[-1])
        for i, text in enumerate(row):
            cells[i].width = widths[i]
            if j % 2 == 1:
                shade_cell(cells[i], "F2F5F9")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_run(p, text, size=8.5, color=INK, bold=(i == 0))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(4)
    add_run(p, "建议阅读路径", size=12, bold=True, color=NAVY)
    paths = [
        ("决策者", "摘要 → 5.6 综合评级 → 7.5 Go/No-Go → 8.3 决策事项"),
        ("临床", "2.2–2.4 声学标志与性能 → 6.1–6.4 研究设计与 PPV → 7.2 前 90 天"),
        ("技术", "3.3–3.5 开源与许可 → 4.3–4.6 架构与模型 → 5.2–5.4 评分与 ADR"),
        ("合规", "6.5–6.7 监管/数据/伦理 → 7.5 阶段门 → 8.4 报告局限"),
    ]
    for role, path in paths:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(4)
        add_run(p, f"{role}：", size=10.5, bold=True, color=NAVY)
        add_run(p, path, size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    add_quote(doc, ["> **一句话：** 先把预期用途收窄，再把数据与证据做扎实；架构是第二位的，但必须第一版就做对。"])
    doc.add_page_break()


# ----------------------------------------------------------------------------- content blocks
def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.42
    p.paragraph_format.space_after = Pt(3)
    add_inline_runs(p, text, size=10.5, color=INK)


def add_caption(doc, text):
    text = text.replace("**", "").strip()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    add_run(p, text, size=9, color=MUTED)


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    p.paragraph_format.line_spacing = 1.38
    p.paragraph_format.space_after = Pt(2)
    add_run(p, "•  ", size=10.5, color=NAVY)
    add_inline_runs(p, text, size=10.5, color=INK)


def add_numbered(doc, num, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(1.5)
    add_run(p, f"{num}. ", size=10.5, bold=True, color=NAVY)
    add_inline_runs(p, text, size=10.5, color=INK)


def add_quote(doc, lines):
    text = " ".join(l[1:].strip() for l in lines)
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Cm(USABLE_CM)
    shade_cell(cell, "EAF2FB")
    set_table_no_borders(tbl)
    # left border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), "1F3A5F")
    tcBorders.append(left)
    tcPr.append(tcBorders)
    set_cell_margins(cell, top=100, start=160, bottom=100, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, text, size=10.5, color=INK)
    # spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(0)
    add_run(sp, " ", size=4)


def add_image(doc, path, alt=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    width = Cm(15.4)
    if "performance" in path.name:
        width = Cm(14.8)
    elif "architecture" in path.name:
        width = Cm(15.0)
    run = p.add_run()
    run.add_picture(str(path), width=width)
    if alt:
        try:
            inline = run._r.xpath(".//wp:inline")[0]
            docPr = inline.xpath("./wp:docPr")[0]
            docPr.set("descr", alt)
        except Exception:
            pass


def add_table(doc, rows):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    tbl = doc.add_table(rows=1, cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_table_borders(tbl)

    # width weights
    weights = []
    for c in range(ncols):
        maxlen = max(len(r[c]) for r in rows)
        weights.append(max(6, min(40, maxlen)))
    total = sum(weights)
    widths_cm = [max(1.2, USABLE_CM * w / total) for w in weights]

    fs = 8.5 if ncols <= 6 else 8.0
    header = rows[0]
    for i, text in enumerate(header):
        cell = tbl.rows[0].cells[i]
        cell.width = Cm(widths_cm[i])
        shade_cell(cell, "1F3A5F")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        add_inline_runs(p, text, size=fs, color=WHITE, base_bold=True)

    set_repeat_table_header(tbl.rows[0])
    for ridx, row in enumerate(rows[1:]):
        cells = tbl.add_row().cells
        prevent_row_split(tbl.rows[-1])
        for i, text in enumerate(row):
            cell = cells[i]
            cell.width = Cm(widths_cm[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if ridx % 2 == 1:
                shade_cell(cell, "F2F5F9")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            add_inline_runs(p, text, size=fs, color=INK)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(0)
    add_run(sp, " ", size=4)


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)
    add_inline_runs(p, text, size=9.0, color=INK)


# ----------------------------------------------------------------------------- markdown parser
def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        line = lines[i].strip()
        if re.match(r"^\|[\s:\-\|]+\|$", line):
            i += 1
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def build():
    text = MD.read_text(encoding="utf-8")
    # split cover from body
    parts = text.split("<!-- pagebreak -->", 1)
    body = parts[1] if len(parts) > 1 else text
    lines = body.splitlines()

    doc = setup_document()
    cover(doc)
    navigation(doc)

    in_references = False
    last_pagebreak = False
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()

        if not line:
            i += 1
            continue

        if line == "<!-- pagebreak -->":
            if not last_pagebreak:
                doc.add_page_break()
                last_pagebreak = True
            i += 1
            continue

        last_pagebreak = False

        # image
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if m:
            alt, rel = m.group(1), m.group(2)
            path = (ROOT / rel).resolve()
            if path.exists():
                add_image(doc, path, alt)
            else:
                add_body(doc, f"[缺失图片：{rel}]")
            i += 1
            continue

        # headings
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            doc.add_heading(title, level=1)
            in_references = title.startswith("参考文献")
            i += 1
            continue

        # table
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        # blockquote
        if line.startswith(">"):
            qlines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                qlines.append(lines[i].strip())
                i += 1
            add_quote(doc, qlines)
            continue

        # caption
        if re.match(r"^\*\*(图|表)\s*\d+", line):
            add_caption(doc, line)
            i += 1
            continue

        # bullet
        if line.startswith("- "):
            add_bullet(doc, line[2:].strip())
            i += 1
            continue

        # numbered
        m = re.match(r"^(\d+)\.\s+(.*)$", line)
        if m:
            add_numbered(doc, m.group(1), m.group(2))
            i += 1
            continue

        # horizontal rule
        if line == "---":
            i += 1
            continue

        # references
        if in_references and re.match(r"^\[\d+\]", line):
            add_reference(doc, line)
            i += 1
            continue

        add_body(doc, line)
        i += 1

    doc.save(OUT)
    print("saved", OUT)
    return OUT


if __name__ == "__main__":
    build()
